import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from office_authoring import MODULES, office_authoring_contract


class OfficeAuthoringContractTests(unittest.TestCase):
    def test_execution_contract_only_probes_for_office_creation(self):
        from agent_execution_harness import execution_contract, execution_policy_for, AgentExecutionPolicy

        with patch("office_authoring.office_authoring_contract", return_value="AUTHORING_HINT") as hint:
            policy = execution_policy_for("Create a PPTX quarterly report")
            self.assertIn("AUTHORING_HINT", execution_contract(policy))
            self.assertEqual(policy, AgentExecutionPolicy.from_public(policy.public()))
            hint.assert_called_once()
        for prompt, kwargs in (
            ("Hello", {}),
            ("Create a downloadable PNG image", {}),
            ("Build an Android APK", {}),
            ("Create a PPTX report, plan only", {}),
            ("Create a PPTX report", {"request_kind": "screen_analysis"}),
            ("Do not generate PPTX. Answer in text.", {}),
        ):
            with self.subTest(prompt=prompt, kwargs=kwargs), \
                    patch("office_authoring.office_authoring_contract") as hint:
                execution_contract(execution_policy_for(prompt, **kwargs))
                hint.assert_not_called()

    def test_missing_modules_do_not_claim_an_available_runtime(self):
        with patch("office_authoring.importlib.util.find_spec", return_value=None):
            self.assertEqual("", office_authoring_contract())

    def test_artifact_quality_rules_do_not_depend_on_installed_office_modules(self):
        from agent_execution_harness import execution_contract, execution_policy_for

        for prompt in ("Create an XLSX inventory workbook", "Create a PPTX budget report",
                       "Export a PDF report", "Create a downloadable PNG column chart"):
            with self.subTest(prompt=prompt), \
                    patch("office_authoring.office_authoring_contract", return_value=""):
                contract = execution_contract(execution_policy_for(prompt))
            self.assertIn("do not infer them from language, locale", contract)
            self.assertIn("neutral numeric formats", contract)
            self.assertIn("including chart axes and cached labels", contract)
            self.assertIn("use a zero baseline", contract)
            self.assertIn("explicit user/template requirement", contract)
            self.assertIn("line, scatter, or logarithmic charts", contract)
            self.assertIn("full chart/drawing bounds inside the print area", contract)
            self.assertIn("Inspect any layout_check issues", contract)
            self.assertIn("Page coverage is not layout approval", contract)
            self.assertIn("saved file and actual preview agree", contract)

    def test_artifact_quality_rules_do_not_turn_read_only_work_into_authoring(self):
        from agent_execution_harness import execution_contract, execution_policy_for

        for prompt, kwargs in (("Hello", {}), ("Create a PPTX report, plan only", {}),
                               ("Analyze the current chart", {"request_kind": "screen_analysis"})):
            with self.subTest(prompt=prompt):
                self.assertNotIn("neutral numeric formats", execution_contract(execution_policy_for(prompt, **kwargs)))

    def test_partial_installation_and_probe_errors_are_safe(self):
        with patch("office_authoring.importlib.util.find_spec",
                   side_effect=[object(), ImportError(), None, ValueError()]), \
                patch("office_authoring.sys.executable", 'C:\\Tools With Space\\python.exe'):
            contract = office_authoring_contract()
        self.assertIn('"C:\\\\Tools With Space\\\\python.exe"', contract)
        self.assertIn("discovered authoring modules: docx.", contract)
        self.assertIn("Discovery is not an import or render check", contract)
        self.assertIn("do not render Office pages", contract)

    def test_complete_runtime_has_formula_and_preview_safeguards(self):
        with patch("office_authoring.importlib.util.find_spec", return_value=object()) as probe:
            contract = office_authoring_contract()
        self.assertEqual([call.args[0] for call in probe.call_args_list], list(MODULES))
        self.assertIn("does not calculate formulas", contract)
        self.assertIn("independently computed cached values", contract)
        self.assertIn("inspect the actual previews", contract)


class OfficeAuthoringRoundTripTests(unittest.TestCase):
    def test_installed_libraries_create_editable_native_files(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        from xlsxwriter import Workbook

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.add_heading("Synthetic quarterly review", 0)
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Item"
            table.cell(0, 1).text = "Amount"
            table.cell(1, 0).text = "A"
            table.cell(1, 1).text = "156"
            doc.save(root / "review.docx")
            self.assertEqual("156", Document(root / "review.docx").tables[0].cell(1, 1).text)

            ppt = Presentation()
            slide = ppt.slides.add_slide(ppt.slide_layouts[5])
            slide.shapes.title.text = "Synthetic data"
            data = CategoryChartData()
            data.categories = ["A", "B", "C"]
            data.add_series("Amount", [156, 184, 180])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   Inches(1), Inches(2), Inches(7), Inches(4), data)
            ppt.save(root / "review.pptx")
            restored = Presentation(root / "review.pptx")
            chart = next(shape.chart for shape in restored.slides[0].shapes if shape.has_chart)
            self.assertEqual((156.0, 184.0, 180.0), chart.series[0].values)
            with zipfile.ZipFile(root / "review.pptx") as archive:
                self.assertTrue(any(name.startswith("ppt/embeddings/") for name in archive.namelist()))

            with Workbook(root / "review.xlsx") as book:
                sheet = book.add_worksheet("Data")
                for row, (qty, price) in enumerate([(13, 12), (23, 8), (18, 10)]):
                    sheet.write_number(row, 0, qty)
                    sheet.write_number(row, 1, price)
                    sheet.write_formula(row, 2, f"=A{row+1}*B{row+1}", None, qty * price)
                sheet.write_formula(3, 2, "=SUM(C1:C3)", None, 520)
            formulas = load_workbook(root / "review.xlsx", data_only=False)
            cached = load_workbook(root / "review.xlsx", data_only=True)
            try:
                self.assertEqual("=SUM(C1:C3)", formulas["Data"]["C4"].value)
                self.assertEqual(520, cached["Data"]["C4"].value)
                self.assertEqual([156, 184, 180], [cached["Data"].cell(row, 3).value for row in (1, 2, 3)])
            finally:
                formulas.close()
                cached.close()


if __name__ == "__main__":
    unittest.main()
