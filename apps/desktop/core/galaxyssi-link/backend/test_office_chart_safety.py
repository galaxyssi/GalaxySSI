from io import BytesIO
import os
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from xml.etree import ElementTree as ET

from office_preview import _validated_source
from office_chart_safety import _validate_workbook
from task_workspace import task_workspace


def repack(raw, replacements):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as before, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as after:
        for name in before.namelist():
            after.writestr(name, replacements.get(name, before.read(name)))
        for name, value in replacements.items():
            if name not in before.namelist():
                after.writestr(name, value)
    return output.getvalue()


class OfficeChartSafetyTests(unittest.TestCase):
    def setUp(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        env = patch.dict(os.environ, {"GALAXYSSI_WORKSPACE_ROOT": temporary.name})
        env.start()
        self.addCleanup(env.stop)
        self.root = task_workspace("chart-preview-test")
        self.source = self.root / "outputs" / "chart.pptx"
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[5])
        data = CategoryChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Amount", [156, 184, 180])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(7), Inches(4), data)
        ppt.save(self.source)
        self.original = self.source.read_bytes()
        with zipfile.ZipFile(self.source) as archive:
            self.embedded = next(n for n in archive.namelist() if n.startswith("ppt/embeddings/"))
            self.workbook = archive.read(self.embedded)
            self.chart = archive.read("ppt/charts/chart1.xml")
            self.rels = archive.read("ppt/charts/_rels/chart1.xml.rels")

    def reject_workbook(self, replacements):
        changed = repack(self.workbook, replacements)
        self.source.write_bytes(repack(self.original, {self.embedded: changed}))
        with self.assertRaises((ValueError, zipfile.BadZipFile, ET.ParseError)):
            _validated_source(self.root, str(self.source))

    def test_passive_native_chart_is_accepted_without_changing_source(self):
        self.assertEqual(self.source, _validated_source(self.root, str(self.source)))
        self.assertEqual(self.original, self.source.read_bytes())

    def test_rejects_macros_nested_packages_connections_and_duplicate_case(self):
        for name in ("xl/vbaProject.bin", "xl/embeddings/nested.xlsx", "xl/connections.xml", "XL/WORKBOOK.XML"):
            with self.subTest(name=name):
                self.reject_workbook({name: b"untrusted"})

    def test_rejects_formulas_even_with_cached_numeric_values(self):
        for formula in ("SUM(1,2)", 'WEBSERVICE("https://example.invalid")'):
            xml = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   '<sheetData><row r="1"><c r="A1"><f>' + formula + '</f><v>3</v></c></row></sheetData></worksheet>')
            with self.subTest(formula=formula):
                self.reject_workbook({"xl/worksheets/sheet1.xml": xml})

    def test_rejects_external_relationships_and_ole(self):
        for extra in (
            b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" TargetMode="External" Target="https://example.invalid"/>',
            b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="worksheets/sheet1.xml"/>',
        ):
            with zipfile.ZipFile(BytesIO(self.workbook)) as book:
                rels = book.read("xl/_rels/workbook.xml.rels").replace(b"</Relationships>", extra + b"</Relationships>")
            with self.subTest(extra=extra):
                self.reject_workbook({"xl/_rels/workbook.xml.rels": rels})

    def test_rejects_macro_workbook_content_type_and_dtd(self):
        with zipfile.ZipFile(BytesIO(self.workbook)) as book:
            types = book.read("[Content_Types].xml")
        self.reject_workbook({"[Content_Types].xml": types.replace(b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")})
        self.reject_workbook({"xl/workbook.xml": b'<!DOCTYPE workbook [<!ENTITY x "value">]><workbook/>'})
        self.reject_workbook({"xl/workbook.xml": '<!DOCTYPE workbook><workbook/>'.encode("utf-16")})

    def test_rejects_unreferenced_workbook_and_ole_reference(self):
        for replacements in (
            {"ppt/embeddings/extra.xlsx": self.workbook},
            {"ppt/charts/_rels/chart1.xml.rels": self.rels.replace(b"relationships/package", b"relationships/oleObject")},
            {"ppt/charts/chart1.xml": self.chart.replace(b'<c:autoUpdate val="0"', b'<c:autoUpdate val="1"')},
            {"ppt/charts/chart1.xml": self.chart.replace(b'r:id="rId1"', b'r:id="rId999"')},
        ):
            self.source.write_bytes(repack(self.original, replacements))
            with self.subTest(parts=list(replacements)), self.assertRaises(ValueError):
                _validated_source(self.root, str(self.source))

    def test_rejects_alternate_encoded_target_to_approved_workbook(self):
        extra = self.rels.replace(b'Id="rId1"', b'Id="rId2"').replace(b"../embeddings/", b"../%65mbeddings/")
        self.source.write_bytes(repack(self.original, {"ppt/charts/_rels/chart2.xml.rels": extra}))
        with self.assertRaises(ValueError):
            _validated_source(self.root, str(self.source))

    def test_rejects_duplicate_relationship_identity(self):
        tree = ET.fromstring(self.rels)
        tree.append(ET.fromstring(ET.tostring(tree[0])))
        self.source.write_bytes(repack(self.original, {"ppt/charts/_rels/chart1.xml.rels": ET.tostring(tree)}))
        with self.assertRaisesRegex(ValueError, "Ambiguous"):
            _validated_source(self.root, str(self.source))

    def test_rejects_expansion_over_budget(self):
        with zipfile.ZipFile(BytesIO(self.workbook)) as book:
            infos = book.infolist()
        infos[0].file_size = 33 * 1024 * 1024
        with patch("office_chart_safety.zipfile.ZipFile.infolist", return_value=infos), self.assertRaises(ValueError):
            _validate_workbook(self.workbook)


if __name__ == "__main__":
    unittest.main()
